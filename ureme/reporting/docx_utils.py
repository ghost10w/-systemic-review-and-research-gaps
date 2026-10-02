"""python-docx helpers for consistent, professional formatting."""

from __future__ import annotations

import datetime as dt
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x59, 0x59, 0x59)
HEADER_FILL = "1F3864"
BAND_FILL = "EAF0F8"


def new_document(font: str = "Calibri", size: int = 11) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.3)
    st = doc.styles["Normal"]
    st.font.name = font
    st.font.size = Pt(size)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.15
    for lvl, sz in ((1, 16), (2, 13), (3, 11.5)):
        h = doc.styles[f"Heading {lvl}"]
        h.font.name = font
        h.font.size = Pt(sz)
        h.font.color.rgb = NAVY
        h.font.bold = True
        h.element.rPr.rFonts.set(qn("w:eastAsia"), font)
        h.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
        h.paragraph_format.space_after = Pt(6)
    _add_page_numbers(doc)
    return doc


def _add_page_numbers(doc: Document) -> None:
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.size = Pt(9)
    run.font.color.rgb = GREY


def title_page(doc: Document, title: str, subtitle: str = "", meta: list[tuple[str, str]] | None = None) -> None:
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = NAVY
    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(subtitle)
        r.font.size = Pt(13)
        r.font.color.rgb = GREY
    doc.add_paragraph()
    meta = meta or []
    meta.append(("Generated", dt.datetime.now().strftime("%d %B %Y, %H:%M")))
    for k, v in meta:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        a = p.add_run(f"{k}: ")
        a.bold = True
        p.add_run(str(v))
    page_break(doc)


def page_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def para(doc: Document, text: str = "", bold_prefix: str | None = None, italic: bool = False, size: float | None = None,
         style: str | None = None, color: RGBColor | None = None):
    p = doc.add_paragraph(style=style)
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    r = p.add_run(text)
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    return p


def bullets(doc: Document, items, numbered: bool = False) -> None:
    for it in items:
        if isinstance(it, tuple):
            para(doc, it[1], bold_prefix=it[0], style="List Number" if numbered else "List Bullet")
        else:
            doc.add_paragraph(str(it), style="List Number" if numbered else "List Bullet")


def callout(doc: Document, text: str, label: str = "Note") -> None:
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    shade(c, "FFF4E5")
    c.paragraphs[0].add_run(f"{label}: ").bold = True
    c.paragraphs[0].add_run(text)
    doc.add_paragraph()


def shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def table(doc: Document, header: list[str], rows: list[list], widths_cm: list[float] | None = None,
          font_size: float = 9, caption: str | None = None, fills: dict[tuple[int, int], str] | None = None):
    if caption:
        p = para(doc, caption, size=10)
        p.runs[0].bold = True
        p.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(str(h))
        r.bold = True
        r.font.size = Pt(font_size)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, HEADER_FILL)
    _repeat_header(t.rows[0])
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, v in enumerate(row):
            cells[ci].text = ""
            r = cells[ci].paragraphs[0].add_run("" if v is None else str(v))
            r.font.size = Pt(font_size)
            if fills and (ri, ci) in fills:
                shade(cells[ci], fills[(ri, ci)])
            elif ri % 2 == 1:
                shade(cells[ci], BAND_FILL)
    if widths_cm:
        for row in t.rows:
            for i, w in enumerate(widths_cm):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def _repeat_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def figure(doc: Document, path: str | Path, caption: str, width_cm: float = 16) -> None:
    if not path or not Path(path).exists():
        return
    doc.add_picture(str(path), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = para(doc, caption, italic=True, size=9)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def toc(doc: Document) -> None:
    """Insert a Word table-of-contents field (updates when the document is opened / F9)."""
    para(doc, "Table of Contents", size=14).runs[0].bold = True
    p = doc.add_paragraph()
    run = p.add_run()
    for kind, text in (("begin", None), (None, 'TOC \\o "1-3" \\h \\z \\u'), ("separate", None), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            if kind == "begin":
                el.set(qn("w:dirty"), "true")
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
        if kind == "separate":
            r2 = OxmlElement("w:t")
            r2.text = "Right-click and choose 'Update field' to build the table of contents."
            run._r.append(r2)
    page_break(doc)


def landscape_section(doc: Document):
    sec = doc.add_section()
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    return sec


def portrait_section(doc: Document):
    sec = doc.add_section()
    sec.orientation = WD_ORIENT.PORTRAIT
    if sec.page_width > sec.page_height:
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    return sec


def fmt(x, nd: int = 2) -> str:
    if x is None:
        return "–"
    try:
        return f"{x:.{nd}f}"
    except (TypeError, ValueError):
        return str(x)
